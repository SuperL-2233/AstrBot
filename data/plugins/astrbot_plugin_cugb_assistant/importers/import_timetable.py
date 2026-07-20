from __future__ import annotations

from pathlib import Path

from ..utils.result import make_result


HEADER_MAPPING = {
    "course_name": "course_name",
    "课程名": "course_name",
    "课程": "course_name",
    "start_time": "start_time",
    "开始时间": "start_time",
    "end_time": "end_time",
    "结束时间": "end_time",
    "location": "location",
    "地点": "location",
    "teacher": "teacher",
    "教师": "teacher",
    "weeks": "weeks",
    "周次": "weeks",
    "weekday": "weekday",
    "星期": "weekday",
}

REQUIRED_HEADERS = {"course_name", "start_time", "end_time", "location"}


class TimetableImporter:
    def __init__(self, timetable_path: Path) -> None:
        self.timetable_path = timetable_path

    def import_courses(self) -> dict:
        if not self.timetable_path.exists():
            return make_result(
                False,
                data=[],
                message="课程表文件不存在，请先上传本地课程表后再执行 /课程表刷新。",
                error_code="TIMETABLE_FILE_NOT_FOUND",
            )

        if self.timetable_path.suffix.lower() != ".xlsx":
            return make_result(
                False,
                data=[],
                message="当前仅支持导入 xlsx 课程表文件。",
                error_code="TIMETABLE_FORMAT_UNSUPPORTED",
            )

        try:
            from openpyxl import load_workbook
        except ImportError:
            return make_result(
                False,
                data=[],
                message="未安装 openpyxl，无法导入 xlsx 课程表。",
                error_code="OPENPYXL_MISSING",
            )

        workbook = load_workbook(self.timetable_path, read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return make_result(
                False,
                data=[],
                message="课程表文件为空。",
                error_code="TIMETABLE_EMPTY",
            )

        raw_headers = [str(item).strip() if item is not None else "" for item in rows[0]]
        mapped_headers = [HEADER_MAPPING.get(header, "") for header in raw_headers]
        if not REQUIRED_HEADERS.issubset(set(mapped_headers)):
            return make_result(
                False,
                data=[],
                message="课程表表头无法识别，需要至少包含课程名、开始时间、结束时间、地点。",
                error_code="TIMETABLE_HEADERS_INVALID",
            )

        courses: list[dict] = []
        for row in rows[1:]:
            row_dict = {}
            for idx, header in enumerate(mapped_headers):
                if not header:
                    continue
                value = row[idx] if idx < len(row) else None
                row_dict[header] = "" if value is None else str(value).strip()
            if not row_dict.get("course_name"):
                continue
            courses.append(
                {
                    "course_name": row_dict["course_name"],
                    "start_time": row_dict["start_time"],
                    "end_time": row_dict["end_time"],
                    "location": row_dict["location"],
                    "teacher": row_dict.get("teacher"),
                    "weeks": row_dict.get("weeks"),
                    "weekday": int(row_dict["weekday"]) if row_dict.get("weekday", "").isdigit() else None,
                    "source": "local_timetable",
                    "raw_data": str(row_dict),
                }
            )

        if not courses:
            return make_result(
                False,
                data=[],
                message="课程表中没有可导入的数据行。",
                error_code="TIMETABLE_NO_ROWS",
            )

        return make_result(True, data=courses, message=f"已解析 {len(courses)} 条课程。")
